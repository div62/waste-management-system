import openpyxl
from pathlib import Path
def get_workbook():

    if DATA_FILE.exists():
        workbook = openpyxl.load_workbook(DATA_FILE)
        sheet = workbook.active
    else:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["ID","Vendor Name","City", "Area", "Waste Type" ])
        workbook.save(DATA_FILE)
    return workbook, sheet
def add_vendor(name, city, area, waste_type):
    workbook, sheet = get_workbook()
    vendor_id = 1
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 1).value is not None:
            vendor_id += 1
    sheet.append([  vendor_id,  name,  city,  area,  waste_type ])
    workbook.save(DATA_FILE)
    return vendor_id
def view_vendors():
    workbook, sheet = get_workbook()
    vendors = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        vendors.append(row)
    workbook.close()
    return vendors
